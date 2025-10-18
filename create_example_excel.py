"""
Script to create an example Excel file for predictions
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

def create_example_predictions_file(filename: str = "predictions.xlsx"):
    """Create an example Excel file with predictions"""
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Predictions"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    ws['A1'] = 'Numero'
    ws['B1'] = 'Costume'
    
    for cell in ['A1', 'B1']:
        ws[cell].fill = header_fill
        ws[cell].font = header_font
        ws[cell].alignment = Alignment(horizontal='center', vertical='center')
    
    example_data = [
        (15, 'Pique'),
        (18, 'Coeur'),
        (22, 'Carreau'),
        (25, 'Trèfle'),
        (30, 'Pique'),
        (35, 'Coeur'),
    ]
    
    for idx, (numero, costume) in enumerate(example_data, start=2):
        ws[f'A{idx}'] = numero
        ws[f'B{idx}'] = costume
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 20
    
    wb.save(filename)
    print(f"✅ Fichier Excel créé: {filename}")
    print(f"📊 Contient {len(example_data)} prédictions d'exemple")

if __name__ == '__main__':
    create_example_predictions_file()
